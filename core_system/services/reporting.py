from __future__ import annotations

import os
import pytz
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Optional

from django.conf import settings
from django.core.files.storage import default_storage
from django.db.models import Sum
from django.utils import timezone

from core_system.models import (
    Member,
    Department,
    MonthlyAssessment,
    MemberAssessment,
    MonthlyDues,
    MembershipFee,
    AidSetAside,
    AidTrackingPost,
    Contribution,
    DeathAid,
    FundTransaction,
    MedicalAid,
    OrganizationFundReport,
)
from core_system.services.compliance import (
    dues_compliance_summary,
    member_dues_status,
    member_dues_status_range,
    member_contribution_status,
    dues_overdue_bucket,
)


def _style_header(ws, row, cols):
    from openpyxl.styles import Alignment, Font, PatternFill, Side, Border

    header_fill = PatternFill(start_color="1b5e20", end_color="1b5e20", fill_type="solid")
    header_font = Font(color="ffffff", bold=True, size=11)
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border


def _auto_width(ws, cols):
    for col in range(1, cols + 1):
        max_len = 0
        for row in ws.iter_rows(min_col=col, max_col=col):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[chr(64 + col)].width = max_len + 4


def _fmt(num):
    if num is None:
        return 0.0
    return round(float(num), 2)


REPORT_TYPES = ("weekly", "monthly", "yearly")


def report_type_label(report_type: str | None) -> str:
    """Human-readable label for a report type (defaults to Monthly)."""
    key = (report_type or "").strip().lower()
    if key == "weekly":
        return "Weekly"
    if key == "yearly":
        return "Yearly"
    return "Monthly"


def resolve_report_range(
    report_type: str | None,
    year: int,
    month: int,
    week: int | None = None,
):
    """Return (period_start, period_end, period_label) for a fund report.

    weekly  -> one 7-day window of the chosen month (week 1-5), label 2026-09-W2
    monthly -> the whole month, label 2026-09
    yearly  -> Jan 1 .. Dec 31 of the year, label 2026
    """
    from calendar import monthrange

    key = (report_type or "monthly").strip().lower()
    if key == "yearly":
        return date(year, 1, 1), date(year, 12, 31), str(year)

    last_day = monthrange(year, month)[1]
    month_end = date(year, month, last_day)

    if key == "weekly":
        try:
            week_no = int(week) if week else 1
        except (TypeError, ValueError):
            week_no = 1
        week_no = min(max(week_no, 1), 5)
        start = date(year, month, 1) + timedelta(days=(week_no - 1) * 7)
        if start > month_end:
            start = month_end
        return start, min(month_end, start + timedelta(days=6)), f"{year}-{month:02d}-W{week_no}"

    return date(year, month, 1), month_end, f"{year}-{month:02d}"


def parse_report_period(period: str | None) -> tuple[str, int, int, int | None] | None:
    """Parse a stored report_period back into (report_type, year, month, week)."""
    text = str(period or "").strip()
    if "-W" in text:
        base, _, week_part = text.rpartition("-W")
        parts = base.split("-")
        if len(parts) != 2:
            return None
        try:
            return ("weekly", int(parts[0]), int(parts[1]), int(week_part))
        except (TypeError, ValueError):
            return None
    parts = text.split("-")
    if len(parts) == 1 and len(parts[0]) == 4 and parts[0].isdigit():
        return ("yearly", int(parts[0]), 1, None)
    if len(parts) == 2:
        try:
            return ("monthly", int(parts[0]), int(parts[1]), None)
        except (TypeError, ValueError):
            return None
    return None


def resolve_report_range_from_period(period: str | None):
    """Inverse of resolve_report_range: stored label -> (start, end, label)."""
    parsed = parse_report_period(period)
    if not parsed:
        return None
    report_type, year, month, week = parsed
    return resolve_report_range(report_type, year, month, week)


def months_covered_by(period_start: date, period_end: date) -> list[str]:
    """YYYY-MM keys touched by a date window (dues are stored per covered month)."""
    months: list[str] = []
    cursor = date(period_start.year, period_start.month, 1)
    while cursor <= period_end:
        months.append(f"{cursor.year}-{cursor.month:02d}")
        cursor = date(cursor.year + (cursor.month // 12), (cursor.month % 12) + 1, 1)
    return months or [f"{period_start.year}-{period_start.month:02d}"]


def compliance_periods(report_type: str | None, year: int, month: int) -> list[tuple[int, int]]:
    """(year, month) pairs dues compliance is computed over.

    Weekly and monthly reports stay month-scoped because dues are monthly;
    a yearly report aggregates every month of the year.
    """
    if (report_type or "").strip().lower() == "yearly":
        return [(year, m) for m in range(1, 13)]
    return [(year, month)]


def merge_dept_summaries(summaries: list[list[dict]]) -> list[dict]:
    """Sum per-month department compliance rows into a single period view."""
    if len(summaries) == 1:
        return summaries[0]
    merged: dict[int, dict] = {}
    for rows in summaries:
        for row in rows:
            acc = merged.get(row["department_id"])
            if acc is None:
                acc = {
                    "department_id": row["department_id"],
                    "department_name": row["department_name"],
                    "total_members": 0,
                    "paid_count": 0,
                    "unpaid_count": 0,
                }
                merged[row["department_id"]] = acc
            acc["total_members"] += row["total_members"]
            acc["paid_count"] += row["paid_count"]
            acc["unpaid_count"] += row["unpaid_count"]
    for acc in merged.values():
        total = acc["total_members"]
        acc["percentage"] = round(acc["paid_count"] / total * 100, 1) if total else 0.0
    return sorted(merged.values(), key=lambda r: r["department_name"])


def generate_overall_report(
    year: int | None = None,
    month: int | None = None,
    wb=None,
    report_type: str = "monthly",
    week: int | None = None,
):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    today = timezone.localdate()
    year = year or today.year
    month = month or today.month
    period_label = resolve_report_range(report_type, year, month, week)[2]
    periods = compliance_periods(report_type, year, month)
    aggregated = len(periods) > 1

    if wb is None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Overall Summary"
    else:
        ws = wb.create_sheet("Overall Summary")
    ws.cell(row=1, column=1, value=f"Compliance Report - {period_label}").font = Font(
        bold=True, size=14
    )
    ws.merge_cells("A1:E1")

    headers = ["Metric", "Value"]
    for ci, h in enumerate(headers, 1):
        ws.cell(row=3, column=ci, value=h)
    _style_header(ws, 3, 2)

    dept_summary = merge_dept_summaries([dues_compliance_summary(y, m) for (y, m) in periods])
    total_members = sum(d["total_members"] for d in dept_summary)
    total_paid = sum(d["paid_count"] for d in dept_summary)
    total_unpaid = sum(d["unpaid_count"] for d in dept_summary)
    overall_pct = round(total_paid / total_members * 100, 1) if total_members else 0.0

    members_metric = "Member-Months (Total)" if aggregated else "Total Active Members"
    paid_metric = "Member-Months Paid" if aggregated else "Total Paid"
    unpaid_metric = "Member-Months Unpaid" if aggregated else "Total Unpaid"

    data_rows = [
        ("Period", period_label),
        (members_metric, total_members),
        (paid_metric, total_paid),
        (unpaid_metric, total_unpaid),
        ("Compliance %", f"{overall_pct}%"),
        ("Basis", "Dues compliance is monthly; this report aggregates "
                  f"{len(periods)} month(s)." if aggregated else "Dues compliance for the reported month."),
    ]
    for ri, (metric, val) in enumerate(data_rows, 4):
        ws.cell(row=ri, column=1, value=metric)
        ws.cell(row=ri, column=2, value=val)
    _auto_width(ws, 2)

    # Sheet 2: Dues by Department
    ws2 = wb.create_sheet("Dues by Department")
    headers2 = ["Department", "Total Members", "Paid", "Unpaid", "Compliance %"]
    for ci, h in enumerate(headers2, 1):
        ws2.cell(row=1, column=ci, value=h)
    _style_header(ws2, 1, len(headers2))

    for ri, d in enumerate(dept_summary, 2):
        ws2.cell(row=ri, column=1, value=d["department_name"])
        ws2.cell(row=ri, column=2, value=d["total_members"])
        ws2.cell(row=ri, column=3, value=d["paid_count"])
        ws2.cell(row=ri, column=4, value=d["unpaid_count"])
        ws2.cell(row=ri, column=5, value=f'{d["percentage"]}%')
        pct = d["percentage"]
        green_fill = PatternFill(start_color="d4edda", end_color="d4edda", fill_type="solid")
        yellow_fill = PatternFill(start_color="fff3cd", end_color="fff3cd", fill_type="solid")
        red_fill = PatternFill(start_color="f8d7da", end_color="f8d7da", fill_type="solid")
        fill = green_fill if pct >= 90 else yellow_fill if pct >= 70 else red_fill
        for ci in range(1, len(headers2) + 1):
            ws2.cell(row=ri, column=ci).fill = fill
    _auto_width(ws2, len(headers2))

    # Sheet 3: Contributions by Department
    ws3 = wb.create_sheet("Contributions by Department")
    headers3 = ["Department", "Post", "Aid Type", "Paid", "Unpaid", "Skipped", "Collection %"]
    for ci, h in enumerate(headers3, 1):
        ws3.cell(row=1, column=ci, value=h)
    _style_header(ws3, 1, len(headers3))

    posts = AidTrackingPost.objects.filter(is_active=True)
    ri = 2
    for post in posts:
        from core_system.services.compliance import contribution_compliance_summary

        dept_contribs = contribution_compliance_summary(post.post_id_PK)
        for dc in dept_contribs:
            ws3.cell(row=ri, column=1, value=dc["department_name"])
            ws3.cell(row=ri, column=2, value=str(post.target_month))
            ws3.cell(row=ri, column=3, value=post.aid_type)
            ws3.cell(row=ri, column=4, value=dc["paid_count"])
            ws3.cell(row=ri, column=5, value=dc["unpaid_count"])
            ws3.cell(row=ri, column=6, value=dc["skipped_count"])
            ws3.cell(row=ri, column=7, value=f'{dc["percentage"]}%')
            ri += 1
    _auto_width(ws3, len(headers3))

    return wb


def generate_department_report(
    dept_id: int,
    year: int | None = None,
    month: int | None = None,
    wb=None,
    sheet_name: str | None = None,
    report_type: str = "monthly",
    week: int | None = None,
):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    today = timezone.localdate()
    year = year or today.year
    month = month or today.month
    period_label = resolve_report_range(report_type, year, month, week)[2]
    periods = compliance_periods(report_type, year, month)

    try:
        dept = Department.objects.get(department_id_PK=dept_id)
    except Department.DoesNotExist:
        return None

    members = Member.objects.filter(department_id_FK=dept).exclude(
        membership_status__iexact="retired"
    ).order_by("full_name")

    dues_sheet = sheet_name or "Dues Compliance"
    contrib_sheet = f"{sheet_name} - Contribution Compliance" if sheet_name else "Contribution Compliance"

    if wb is None:
        wb = Workbook()
        ws = wb.active
        ws.title = dues_sheet
    else:
        ws = wb.create_sheet(dues_sheet)
    ws.cell(row=1, column=1, value=f"{dept.name} - Dues Compliance ({period_label})").font = Font(
        bold=True, size=12
    )
    ws.merge_cells("A1:E1")

    headers = ["Member Name", "Employee ID", "Dues Status", "Overdue Bucket"]
    for ci, h in enumerate(headers, 1):
        ws.cell(row=3, column=ci, value=h)
    _style_header(ws, 3, len(headers))

    bucket_period = periods[-1]
    for ri, m in enumerate(members, 4):
        ws.cell(row=ri, column=1, value=m.full_name)
        ws.cell(row=ri, column=2, value=m.employee_id or "")
        status = member_dues_status_range(m, periods)
        ws.cell(row=ri, column=3, value=status)
        ws.cell(row=ri, column=4, value=dues_overdue_bucket(m, *bucket_period) or "")
    _auto_width(ws, len(headers))

    ws2 = wb.create_sheet(contrib_sheet)
    posts = AidTrackingPost.objects.filter(is_active=True)
    headers2 = ["Member Name", "Employee ID", "Post", "Aid Type", "Status"]
    for ci, h in enumerate(headers2, 1):
        ws2.cell(row=1, column=ci, value=h)
    _style_header(ws2, 1, len(headers2))

    ri = 2
    for m in members:
        for post in posts:
            status = member_contribution_status(m, post.post_id_PK)
            ws2.cell(row=ri, column=1, value=m.full_name)
            ws2.cell(row=ri, column=2, value=m.employee_id or "")
            ws2.cell(row=ri, column=3, value=str(post.target_month))
            ws2.cell(row=ri, column=4, value=post.aid_type)
            ws2.cell(row=ri, column=5, value=status)
            ri += 1
    _auto_width(ws2, len(headers2))

    return wb


def generate_contribution_report(post_id: int | None = None, wb=None):
    from openpyxl import Workbook

    if wb is None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Contribution Report"
    else:
        ws = wb.create_sheet("Contribution Report")

    posts = AidTrackingPost.objects.filter(is_active=True)
    if post_id:
        posts = posts.filter(post_id_PK=post_id)

    headers = [
        "Post ID", "Aid Type", "Target Month", "Department",
        "Member Name", "Employee ID", "Status", "Expected Amount",
        "Paid Amount", "Payment Date",
    ]
    for ci, h in enumerate(headers, 1):
        ws.cell(row=1, column=ci, value=h)
    _style_header(ws, 1, len(headers))

    ri = 2
    for post in posts:
        contribs = Contribution.objects.filter(
            aid_tracking_post_id_FK=post,
        ).select_related("member_id_FK__department_id_FK")
        for c in contribs:
            dept_name = (
                c.member_id_FK.department_id_FK.name
                if c.member_id_FK.department_id_FK
                else "N/A"
            )
            ws.cell(row=ri, column=1, value=post.post_id_PK)
            ws.cell(row=ri, column=2, value=post.aid_type)
            ws.cell(row=ri, column=3, value=post.target_month)
            ws.cell(row=ri, column=4, value=dept_name)
            ws.cell(row=ri, column=5, value=c.member_id_FK.full_name)
            ws.cell(row=ri, column=6, value=c.member_id_FK.employee_id or "")
            ws.cell(row=ri, column=7, value=c.status)
            ws.cell(row=ri, column=8, value=_fmt(c.expected_amount))
            ws.cell(row=ri, column=9, value=_fmt(c.paid_amount))
            ws.cell(row=ri, column=10, value=str(c.payment_date) if c.payment_date else "")
            ri += 1
    _auto_width(ws, len(headers))

    return wb


def approve_auditor_report(report_id: int, officer, remarks: str = "") -> dict | None:
    from core_system.models import AuditFindingsReport

    try:
        report = AuditFindingsReport.objects.get(
            audit_report_id_PK=report_id,
            report_status="Submitted",
        )
    except AuditFindingsReport.DoesNotExist:
        return None

    report.report_status = "Approved"
    report.certification_status = "Certified"
    report.certified_by_user_id_FK = officer
    report.save(update_fields=["report_status", "certification_status", "certified_by_user_id_FK"])

    return {
        "report_id": report.audit_report_id_PK,
        "status": report.report_status,
        "certification_status": report.certification_status,
    }


def request_report_revision(report_id: int, officer, remarks: str = "") -> dict | None:
    from core_system.models import AuditFindingsReport

    try:
        report = AuditFindingsReport.objects.get(
            audit_report_id_PK=report_id,
            report_status="Submitted",
        )
    except AuditFindingsReport.DoesNotExist:
        return None

    report.report_status = "Revision Requested"
    report.findings_summary += f"\n\n--- Revision requested by {officer.full_name}: {remarks}"
    report.save(update_fields=["report_status", "findings_summary"])

    return {
        "report_id": report.audit_report_id_PK,
        "status": report.report_status,
    }


def create_auditor_report(year: int, month: int, officer) -> dict:
    from core_system.models import AuditFindingsReport

    dept_summary = dues_compliance_summary(year, month)
    total_members = sum(d["total_members"] for d in dept_summary)
    total_paid = sum(d["paid_count"] for d in dept_summary)
    total_unpaid = sum(d["unpaid_count"] for d in dept_summary)
    pct = round(total_paid / total_members * 100, 1) if total_members else 0.0

    period = f"{year}-{month:02d}"
    low_depts = [d for d in dept_summary if d["percentage"] < 70]

    findings = (
        f"Compliance report for period {period}.\n"
        f"Overall compliance: {pct}%\n"
        f"Total active members: {total_members}\n"
        f"Paid: {total_paid}, Unpaid: {total_unpaid}\n"
    )
    if low_depts:
        findings += f"Departments below 70% threshold: {', '.join(d['department_name'] for d in low_depts)}\n"

    report = AuditFindingsReport.objects.create(
        report_title=f"Compliance Report - {period}",
        report_period=period,
        findings_summary=findings,
        report_status="Draft",
        prepared_by_user_id_FK=officer,
        prepared_date=timezone.localdate(),
        presentation_status="Draft",
        certification_status="Pending",
    )

    return {
        "report_id": report.audit_report_id_PK,
        "title": report.report_title,
        "period": report.report_period,
        "findings_summary": report.findings_summary,
        "status": report.report_status,
        "prepared_date": str(report.prepared_date),
    }


# Fund outflow sources that represent an aid disbursement: the claim itself
# (normal release) or the tracking post (release paid straight from the fund).
AID_PAYOUT_SOURCE_TYPES = ("medical_aid", "death_aid", "aid_post_payment")

_AID_TYPE_LABELS = {"medical_aid": "Medical Aid", "death_aid": "Death Aid"}


def _aid_type_label(raw) -> str:
    """Display label for a stored aid type (post, set-aside or claim value)."""
    key = str(raw or "").strip()
    if not key:
        return ""
    return _AID_TYPE_LABELS.get(key, key.replace("_", " ").title())


def _aid_disbursement_label(value) -> str:
    if value == "fund":
        return "Fund"
    if value == "direct":
        return "Direct"
    return ""


def build_aid_payout_rows(start_dt, end_dt) -> list[dict]:
    """Aid disbursements booked to the fund inside ``start_dt..end_dt``.

    One row per outflow transaction, joined back to the claim it paid so the
    report shows what was APPROVED next to what was actually DISBURSED —
    Treasurer override amounts and general-fund top-ups stop being invisible.
    """
    transactions = list(
        FundTransaction.objects.filter(
            direction="outflow",
            source_type__in=AID_PAYOUT_SOURCE_TYPES,
            recorded_at__gte=start_dt,
            recorded_at__lte=end_dt,
        ).select_related("recorded_by_user_id_FK").order_by("recorded_at")
    )
    medical = {
        m.medical_aid_id_PK: m
        for m in MedicalAid.objects.select_related("member_id_FK", "released_by_user_id_FK")
    }
    death = {
        d.death_aid_id_PK: d
        for d in DeathAid.objects.select_related(
            "claimant_id_FK", "member_id_FK", "released_by_user_id_FK"
        )
    }
    posts = {
        p.post_id_PK: p
        for p in AidTrackingPost.objects.select_related("archive_id_FK")
    }

    def claim_context(source_type, source_id):
        if source_type == "medical_aid":
            claim = medical.get(source_id)
            if claim is None:
                return {}
            return {
                "aid_type": "medical_aid",
                "reference": f"MA-{claim.medical_aid_id_PK}",
                "recipient": claim.member_id_FK.full_name if claim.member_id_FK else "",
                "approved": float(claim.validated_aid_amount or 0),
                "disbursement": _aid_disbursement_label(claim.disbursement_source),
                "release_reference": claim.release_reference or "",
                "released_by": (
                    claim.released_by_user_id_FK.full_name
                    if claim.released_by_user_id_FK else ""
                ),
                "status": claim.status or "",
            }
        if source_type == "death_aid":
            claim = death.get(source_id)
            if claim is None:
                return {}
            return {
                "aid_type": "death_aid",
                "reference": f"DA-{claim.death_aid_id_PK}",
                "recipient": claim.deceased_name or "",
                "approved": float(claim.benefit_amount or 0),
                "disbursement": _aid_disbursement_label(claim.disbursement_source),
                "release_reference": claim.release_reference or "",
                "released_by": (
                    claim.released_by_user_id_FK.full_name
                    if claim.released_by_user_id_FK else ""
                ),
                "status": claim.status or "",
            }
        return {}

    rows = []
    for t in transactions:
        ctx = claim_context(t.source_type, t.source_id)
        if t.source_type == "aid_post_payment":
            post = posts.get(t.source_id)
            if post is not None:
                ctx = {
                    "aid_type": post.aid_type,
                    "reference": f"POST-{post.post_id_PK}",
                    "recipient": post.archive_id_FK.member_name if post.archive_id_FK else "",
                    "approved": float(post.total_expected or 0),
                    "disbursement": "",
                    "release_reference": "",
                    "released_by": "",
                    "status": post.finish_status or post.status or "",
                }
                claim_ctx = claim_context(post.source_type, post.source_id)
                if claim_ctx:
                    ctx["disbursement"] = claim_ctx["disbursement"]
                    ctx["release_reference"] = claim_ctx["release_reference"]
                    ctx["released_by"] = claim_ctx["released_by"]
                    if not ctx["recipient"]:
                        ctx["recipient"] = claim_ctx["recipient"]

        paid = float(t.amount)
        approved = ctx.get("approved")
        recorded_by = t.recorded_by_user_id_FK.full_name if t.recorded_by_user_id_FK else ""
        rows.append({
            "date": str(t.recorded_at.date()) if t.recorded_at else "",
            "aid_type": _aid_type_label(ctx.get("aid_type")) or _aid_type_label(t.source_type),
            "reference": ctx.get("reference") or f"{t.source_type} #{t.source_id}",
            "recipient": ctx.get("recipient") or "",
            "approved": approved,
            "paid": paid,
            "variance": (approved - paid) if approved is not None else None,
            "disbursement": ctx.get("disbursement") or "-",
            "release_reference": ctx.get("release_reference") or "-",
            "released_by": ctx.get("released_by") or recorded_by or "-",
            "status": ctx.get("status") or "-",
        })
    return rows


def build_aid_setaside_rows() -> tuple[list[dict], list[dict]]:
    """``(summary_rows, detail_rows)`` for the aid set-aside balance sheet.

    Balances are cumulative to the moment the report is produced rather than
    cut at the period end: ``AidSetAside.amount_released`` is a running total
    with no release date, so a period cut-off would make
    "collected - released" disagree with itself. Period-scoped disbursements
    live in :func:`build_aid_payout_rows`.

    "Paid from General Fund" is the shortfall — a claim drawn down for more
    than its own members' earmark.
    """
    posts = list(
        AidTrackingPost.objects.select_related("archive_id_FK")
        .order_by("aid_type", "target_month", "post_id_PK")
    )

    linked = {
        row["aid_tracking_post_id_FK_id"]: row
        for row in AidSetAside.objects.filter(aid_tracking_post_id_FK__isnull=False)
        .values("aid_tracking_post_id_FK_id")
        .annotate(collected=Sum("amount"), released=Sum("amount_released"))
    }
    pooled = AidSetAside.objects.filter(aid_tracking_post_id_FK__isnull=True).aggregate(
        collected=Sum("amount"), released=Sum("amount_released")
    )
    pooled_collected = float(pooled["collected"] or 0)
    pooled_released = float(pooled["released"] or 0)

    claim_to_post = {
        (p.source_type, p.source_id): p.post_id_PK
        for p in posts
        if p.source_type and p.source_id
    }
    paid_by_post: dict[int, float] = {}
    for t in FundTransaction.objects.filter(
        direction="outflow", source_type__in=AID_PAYOUT_SOURCE_TYPES
    ):
        post_id = (
            t.source_id
            if t.source_type == "aid_post_payment"
            else claim_to_post.get((t.source_type, t.source_id))
        )
        if post_id:
            paid_by_post[post_id] = paid_by_post.get(post_id, 0.0) + float(t.amount)

    detail_rows = []
    for p in posts:
        agg = linked.get(p.post_id_PK) or {}
        collected = float(agg.get("collected") or 0)
        released = float(agg.get("released") or 0)
        expected = float(p.total_expected or 0)
        contributions = float(p.total_collected or 0)
        paid = paid_by_post.get(p.post_id_PK, 0.0)
        detail_rows.append({
            "post_id": p.post_id_PK,
            "aid_type": _aid_type_label(p.aid_type),
            "aid_type_key": str(p.aid_type or ""),
            "target_month": p.target_month or "",
            "expected": expected,
            "contributions": contributions,
            "setaside_collected": collected,
            "setaside_released": released,
            "remaining": collected - released,
            "paid": paid,
            "from_general_fund": max(0.0, paid - released),
            "members_still_owe": max(0.0, expected - contributions),
            "status": p.finish_status or p.status or "",
        })

    def summarise(label: str, rows: list[dict]) -> dict:
        return {
            "aid_type": label,
            "posts": len(rows),
            "setaside_collected": sum(r["setaside_collected"] for r in rows),
            "setaside_released": sum(r["setaside_released"] for r in rows),
            "remaining": sum(r["remaining"] for r in rows),
            "from_general_fund": sum(r["from_general_fund"] for r in rows),
            "members_still_owe": sum(r["members_still_owe"] for r in rows),
        }

    summary_rows = []
    for key, label in _AID_TYPE_LABELS.items():
        summary_rows.append(summarise(label, [r for r in detail_rows if r["aid_type_key"] == key]))
    others = [
        r for r in detail_rows if r["aid_type_key"] not in _AID_TYPE_LABELS
    ]
    if others:
        summary_rows.append(summarise("Other", others))
    if pooled_collected or pooled_released:
        summary_rows.append({
            "aid_type": "Pooled (no claim yet)",
            "posts": 0,
            "setaside_collected": pooled_collected,
            "setaside_released": pooled_released,
            "remaining": pooled_collected - pooled_released,
            "from_general_fund": 0.0,
            "members_still_owe": 0.0,
        })

    total = summarise("TOTAL", summary_rows)
    total["posts"] = len(detail_rows)
    summary_rows.append(total)
    return summary_rows, detail_rows


def generate_organization_fund_report(
    year: int,
    month: int,
    report_type: str = "monthly",
    wb=None,
    week: int | None = None,
):
    from openpyxl import Workbook
    from openpyxl.styles import Font

    if wb is None:
        wb = Workbook()
        ws = wb.active
        ws.title = "Fund Summary"
    else:
        ws = wb.create_sheet("Fund Summary")

    period_start, period_end, period_label = resolve_report_range(report_type, year, month, week)
    covered_months = months_covered_by(period_start, period_end)

    start_dt = pytz.UTC.localize(datetime.combine(period_start, datetime.min.time()))
    end_dt = pytz.UTC.localize(datetime.combine(period_end, datetime.max.time().replace(microsecond=0)))

    transactions = FundTransaction.objects.filter(
        recorded_at__gte=start_dt,
        recorded_at__lte=end_dt,
    ).select_related("recorded_by_user_id_FK").order_by("recorded_at")

    inflows = [t for t in transactions if t.direction == "inflow"]
    outflows = [t for t in transactions if t.direction == "outflow"]

    total_inflow = sum(float(t.amount) for t in inflows)
    total_outflow = sum(float(t.amount) for t in outflows)
    net = total_inflow - total_outflow

    contributions = Contribution.objects.filter(
        payment_date__gte=period_start,
        payment_date__lte=period_end,
        status__in=["PAID", "RECORDED", "PENDING_VERIFICATION"],
    ).select_related("member_id_FK", "aid_tracking_post_id_FK").order_by("payment_date")

    total_contributions = sum(float(c.paid_amount) for c in contributions)

    # Monthly dues matched by the month they COVER (YYYY-MM), regardless of
    # when the fund_transaction entry was recorded — keeps the workbook
    # consistent with the treasurer's Monthly Dues Summary for the period.
    # A weekly report covers the month containing that week; a yearly report
    # covers every month of the year.
    dues = MonthlyDues.objects.filter(
        month_covered__in=covered_months,
    ).exclude(payment_status__iexact="rejected").select_related("member_id_FK").order_by("member_id_FK__full_name")
    total_dues = sum(float(d.amount) for d in dues)

    # Sheet 1: Summary — reuse the sheet created above ("Fund Summary").
    # Never use wb.active here: in the unified workbook it points to the first
    # sheet (Overall Summary) and would overwrite its content.
    ws.cell(row=1, column=1, value=f"Organization Fund Report - {period_label}").font = Font(bold=True, size=14)
    ws.merge_cells("A1:D1")

    headers = ["Metric", "Value"]
    for ci, h in enumerate(headers, 1):
        ws.cell(row=3, column=ci, value=h)
    _style_header(ws, 3, 2)

    dues_metric = (
        "Monthly Dues Collected (covered month)"
        if len(covered_months) == 1
        else "Monthly Dues Collected (covered months)"
    )
    month_starts = [
        date(int(ym[:4]), int(ym[5:7]), 1) for ym in covered_months
    ]
    assessments = MonthlyAssessment.objects.filter(month__in=month_starts).exclude(
        status__in=[MonthlyAssessment.STATUS_REJECTED, MonthlyAssessment.STATUS_RETURNED]
    )
    outstanding_rows = list(
        MemberAssessment.objects.filter(assessment_id_FK__in=assessments)
        .select_related("member_id_FK", "assessment_id_FK")
        .order_by("member_id_FK__full_name", "assessment_id_FK__month")
    )
    total_outstanding = sum(float(r.outstanding_balance) for r in outstanding_rows)

    dues_deposits = (
        MonthlyAssessment.objects.filter(
            deposited_at__gte=start_dt, deposited_at__lte=end_dt
        )
        .exclude(deposit_reference="")
        .exclude(deposit_reference__isnull=True)
        .select_related("deposited_by_id_FK")
    )
    cash_movements = [
        {
            "date": t.recorded_at,
            "kind": "Deposit" if t.direction == "inflow" else "Withdrawal",
            "reference": t.reference_number or "",
            "description": t.description or "",
            "amount": float(t.amount),
            "recorded_by": t.recorded_by_user_id_FK.full_name if t.recorded_by_user_id_FK else "",
        }
        for t in transactions if t.source_type == "other_transaction"
    ] + [
        {
            "date": a.deposited_at,
            "kind": "Dues Deposit",
            "reference": a.deposit_reference or "",
            "description": f"Monthly Dues - {a.month_label}",
            "amount": float(a.deposited_amount or 0),
            "recorded_by": a.deposited_by_id_FK.full_name if a.deposited_by_id_FK else "",
        }
        for a in dues_deposits
    ]
    cash_movements.sort(key=lambda r: (r["date"] is None, r["date"]))

    # Sheet 8/9 source data: what was actually disbursed for aid claims in the
    # period, and the running set-aside position behind it.
    aid_payouts = build_aid_payout_rows(start_dt, end_dt)
    total_aid_payout = sum(r["paid"] for r in aid_payouts)
    aid_setaside_summary, aid_setaside_detail = build_aid_setaside_rows()

    data_rows = [
        ("Period", period_label),
        ("Report Type", report_type_label(report_type)),
        ("Date Range", f"{period_start} to {period_end}"),
        ("Total Inflows", total_inflow),
        ("Total Outflows", total_outflow),
        ("Net Fund Position", net),
        ("Total Contributions Collected", total_contributions),
        (dues_metric, total_dues),
        ("Total Outstanding Balance", total_outstanding),
        ("Deposits and Withdrawals", len(cash_movements)),
        ("Aid Disbursed This Period", total_aid_payout),
        ("Aid Set-Aside Reserve Remaining", aid_setaside_summary[-1]["remaining"]),
        ("Number of Transactions", len(transactions)),
    ]
    for ri, (metric, val) in enumerate(data_rows, 4):
        ws.cell(row=ri, column=1, value=metric)
        ws.cell(row=ri, column=2, value=val)
    _auto_width(ws, 2)

    # Sheet 2: Inflows
    ws2 = wb.create_sheet("Inflows")
    headers2 = ["Date", "Source Type", "Source ID", "Amount", "Description", "Recorded By"]
    for ci, h in enumerate(headers2, 1):
        ws2.cell(row=1, column=ci, value=h)
    _style_header(ws2, 1, len(headers2))
    for ri, t in enumerate(inflows, 2):
        ws2.cell(row=ri, column=1, value=str(t.recorded_at.date()) if t.recorded_at else "")
        ws2.cell(row=ri, column=2, value=t.get_source_type_display() or "")
        ws2.cell(row=ri, column=3, value=t.source_id or "")
        ws2.cell(row=ri, column=4, value=float(t.amount))
        ws2.cell(row=ri, column=5, value=t.description or "")
        ws2.cell(row=ri, column=6, value=t.recorded_by_user_id_FK.full_name if t.recorded_by_user_id_FK else "")
    _auto_width(ws2, len(headers2))

    # Sheet 3: Outflows
    ws3 = wb.create_sheet("Outflows")
    headers3 = ["Date", "Source Type", "Source ID", "Amount", "Description", "Recorded By"]
    for ci, h in enumerate(headers3, 1):
        ws3.cell(row=1, column=ci, value=h)
    _style_header(ws3, 1, len(headers3))
    for ri, t in enumerate(outflows, 2):
        ws3.cell(row=ri, column=1, value=str(t.recorded_at.date()) if t.recorded_at else "")
        ws3.cell(row=ri, column=2, value=t.get_source_type_display() or "")
        ws3.cell(row=ri, column=3, value=t.source_id or "")
        ws3.cell(row=ri, column=4, value=float(t.amount))
        ws3.cell(row=ri, column=5, value=t.description or "")
        ws3.cell(row=ri, column=6, value=t.recorded_by_user_id_FK.full_name if t.recorded_by_user_id_FK else "")
    _auto_width(ws3, len(headers3))

    # Sheet 4: Contributions
    ws4 = wb.create_sheet("Contributions")
    headers4 = ["Date", "Member Name", "Employee ID", "Department", "Post ID", "Aid Type", "Expected", "Paid", "Status"]
    for ci, h in enumerate(headers4, 1):
        ws4.cell(row=1, column=ci, value=h)
    _style_header(ws4, 1, len(headers4))
    for ri, c in enumerate(contributions, 2):
        member = c.member_id_FK
        post = c.aid_tracking_post_id_FK
        ws4.cell(row=ri, column=1, value=str(c.payment_date) if c.payment_date else "")
        ws4.cell(row=ri, column=2, value=member.full_name if member else "")
        ws4.cell(row=ri, column=3, value=member.employee_id or "" if member else "")
        ws4.cell(row=ri, column=4, value=member.department or "" if member else "")
        ws4.cell(row=ri, column=5, value=post.post_id_PK if post else "")
        ws4.cell(row=ri, column=6, value=post.aid_type if post else "")
        ws4.cell(row=ri, column=7, value=float(c.expected_amount))
        ws4.cell(row=ri, column=8, value=float(c.paid_amount))
        ws4.cell(row=ri, column=9, value=c.status)
    _auto_width(ws4, len(headers4))

    # Sheet 5: Monthly Dues (by covered month)
    ws5 = wb.create_sheet("Monthly Dues")
    headers5 = ["Member Name", "Month Covered", "Amount", "Payment Method", "Payment Status"]
    for ci, h in enumerate(headers5, 1):
        ws5.cell(row=1, column=ci, value=h)
    _style_header(ws5, 1, len(headers5))
    for ri, d in enumerate(dues, 2):
        ws5.cell(row=ri, column=1, value=d.member_id_FK.full_name if d.member_id_FK else "")
        ws5.cell(row=ri, column=2, value=d.month_covered or "")
        ws5.cell(row=ri, column=3, value=float(d.amount))
        ws5.cell(row=ri, column=4, value=d.payment_method or "")
        ws5.cell(row=ri, column=5, value=d.payment_status or "")
    _auto_width(ws5, len(headers5))

    # Sheet 6: Deposit & Withdrawal history — cash movements the Treasurer
    # actually touched at the bank, kept apart from raw fund inflows/outflows.
    ws6 = wb.create_sheet("Deposits and Withdrawals")
    headers6 = ["Date", "Type", "Reference", "Description", "Amount", "Recorded By"]
    for ci, h in enumerate(headers6, 1):
        ws6.cell(row=1, column=ci, value=h)
    _style_header(ws6, 1, len(headers6))
    for ri, mv in enumerate(cash_movements, 2):
        ws6.cell(row=ri, column=1, value=str(mv["date"].date()) if mv["date"] else "")
        ws6.cell(row=ri, column=2, value=mv["kind"])
        ws6.cell(row=ri, column=3, value=mv["reference"])
        ws6.cell(row=ri, column=4, value=mv["description"])
        ws6.cell(row=ri, column=5, value=mv["amount"])
        ws6.cell(row=ri, column=6, value=mv["recorded_by"])
    _auto_width(ws6, len(headers6))

    # Sheet 7: Outstanding balances for the covered months.
    ws7 = wb.create_sheet("Outstanding Balances")
    headers7 = ["Member Name", "Employee ID", "Department", "Month Covered",
                "Expected", "Recorded", "Outstanding", "Status"]
    for ci, h in enumerate(headers7, 1):
        ws7.cell(row=1, column=ci, value=h)
    _style_header(ws7, 1, len(headers7))
    for ri, r in enumerate(outstanding_rows, 2):
        member = r.member_id_FK
        ws7.cell(row=ri, column=1, value=member.full_name if member else "")
        ws7.cell(row=ri, column=2, value=member.employee_id if member else "")
        ws7.cell(row=ri, column=3, value=member.department if member else "")
        ws7.cell(row=ri, column=4, value=r.assessment_id_FK.month_label)
        ws7.cell(row=ri, column=5, value=float(r.standard_assessment))
        ws7.cell(row=ri, column=6, value=float(r.actual_deduction))
        ws7.cell(row=ri, column=7, value=float(r.outstanding_balance))
        ws7.cell(row=ri, column=8, value=r.get_status_display())
    total_row = len(outstanding_rows) + 2
    ws7.cell(row=total_row, column=1, value="TOTAL OUTSTANDING").font = Font(bold=True)
    ws7.cell(row=total_row, column=7, value=total_outstanding).font = Font(bold=True)
    _auto_width(ws7, len(headers7))

    # Sheet 8: Aid payouts — every fund disbursement for a claim in the period,
    # showing the approved amount next to what was actually paid.
    ws8 = wb.create_sheet("Aid Payouts by Claim")
    headers8 = ["Date", "Aid Type", "Reference", "Recipient / Claimant", "Approved",
                "Disbursed", "Variance", "Disbursement", "Release Reference",
                "Released By", "Status"]
    for ci, h in enumerate(headers8, 1):
        ws8.cell(row=1, column=ci, value=h)
    _style_header(ws8, 1, len(headers8))
    for ri, r in enumerate(aid_payouts, 2):
        ws8.cell(row=ri, column=1, value=r["date"])
        ws8.cell(row=ri, column=2, value=r["aid_type"])
        ws8.cell(row=ri, column=3, value=r["reference"])
        ws8.cell(row=ri, column=4, value=r["recipient"])
        ws8.cell(row=ri, column=5, value=r["approved"] if r["approved"] is not None else "")
        ws8.cell(row=ri, column=6, value=r["paid"])
        ws8.cell(row=ri, column=7, value=r["variance"] if r["variance"] is not None else "")
        ws8.cell(row=ri, column=8, value=r["disbursement"])
        ws8.cell(row=ri, column=9, value=r["release_reference"])
        ws8.cell(row=ri, column=10, value=r["released_by"])
        ws8.cell(row=ri, column=11, value=r["status"])
    payout_total_row = len(aid_payouts) + 2
    ws8.cell(row=payout_total_row, column=4, value="TOTAL DISBURSED").font = Font(bold=True)
    ws8.cell(row=payout_total_row, column=6, value=total_aid_payout).font = Font(bold=True)
    _auto_width(ws8, len(headers8))

    # Sheet 9: Aid set-aside balance — per-aid-type reserve on top, per-claim
    # detail underneath. Balances are cumulative to report date (see
    # build_aid_setaside_rows); "Paid from General Fund" is the shortfall the
    # earmark could not cover.
    ws9 = wb.create_sheet("Aid Set-Aside Balance")
    ws9.cell(row=1, column=1, value=f"Aid Set-Aside Balance - {period_label}").font = Font(bold=True, size=14)
    summary_headers = ["Aid Type", "Posts", "Set-Aside Collected", "Set-Aside Released",
                       "Remaining Reserve", "Paid from General Fund", "Members Still Owe"]
    for ci, h in enumerate(summary_headers, 1):
        ws9.cell(row=3, column=ci, value=h)
    _style_header(ws9, 3, len(summary_headers))
    for offset, r in enumerate(aid_setaside_summary, 4):
        ws9.cell(row=offset, column=1, value=r["aid_type"])
        ws9.cell(row=offset, column=2, value=r["posts"])
        ws9.cell(row=offset, column=3, value=r["setaside_collected"])
        ws9.cell(row=offset, column=4, value=r["setaside_released"])
        ws9.cell(row=offset, column=5, value=r["remaining"])
        ws9.cell(row=offset, column=6, value=r["from_general_fund"])
        ws9.cell(row=offset, column=7, value=r["members_still_owe"])
        if r["aid_type"] == "TOTAL":
            for ci in range(1, len(summary_headers) + 1):
                ws9.cell(row=offset, column=ci).font = Font(bold=True)

    detail_title_row = 4 + len(aid_setaside_summary) + 1
    ws9.cell(row=detail_title_row, column=1, value="Claim Detail").font = Font(bold=True, size=12)
    detail_header_row = detail_title_row + 1
    detail_headers = ["Post ID", "Aid Type", "Target Month", "Expected", "Contributions",
                      "Set-Aside Collected", "Set-Aside Released", "Remaining",
                      "Paid to Claim", "Paid from General Fund", "Members Still Owe", "Status"]
    for ci, h in enumerate(detail_headers, 1):
        ws9.cell(row=detail_header_row, column=ci, value=h)
    _style_header(ws9, detail_header_row, len(detail_headers))
    for i, r in enumerate(aid_setaside_detail):
        row = detail_header_row + 1 + i
        values = [r["post_id"], r["aid_type"], r["target_month"], r["expected"],
                  r["contributions"], r["setaside_collected"], r["setaside_released"],
                  r["remaining"], r["paid"], r["from_general_fund"],
                  r["members_still_owe"], r["status"]]
        for ci, value in enumerate(values, 1):
            ws9.cell(row=row, column=ci, value=value)
    _auto_width(ws9, len(detail_headers))

    return wb


def create_organization_fund_report(
    officer,
    year: int,
    month: int,
    report_type: str = "monthly",
    week: int | None = None,
) -> dict:
    period = resolve_report_range(report_type, year, month, week)[2]

    wb = generate_organization_fund_report(year, month, report_type, week=week)

    filename = f"fund_report_{report_type}_{period}.xlsx"
    file_path = f"reports/{filename}"
    default_storage_path = getattr(settings, "MEDIA_ROOT", "")
    full_path = os.path.join(str(default_storage_path), file_path) if default_storage_path else file_path

    os.makedirs(os.path.dirname(full_path), exist_ok=True)
    wb.save(full_path)

    report = OrganizationFundReport.objects.create(
        report_period=period,
        report_type=report_type,
        report_status="Draft",
        file_path=file_path,
        prepared_by_user_id_FK=officer,
    )

    return {
        "report_id": report.report_id_PK,
        "period": report.report_period,
        "report_type": report.report_type,
        "status": report.report_status,
        "file_path": report.file_path,
    }


def generate_unified_report(
    year: int | None = None,
    month: int | None = None,
    sections: list[str] | None = None,
    report_type: str = "monthly",
    week: int | None = None,
) -> Workbook:
    from openpyxl import Workbook

    today = timezone.localdate()
    year = year or today.year
    month = month or today.month

    if sections is None:
        sections = ["overall", "department", "contribution", "fund"]

    wb = Workbook()
    default_sheet = wb.active
    wb.remove(default_sheet)

    if "overall" in sections:
        generate_overall_report(year, month, wb, report_type=report_type, week=week)

    if "department" in sections:
        departments = Department.objects.filter(is_active=True).order_by("name")
        for dept in departments:
            generate_department_report(
                dept.department_id_PK,
                year,
                month,
                wb,
                sheet_name=dept.name,
                report_type=report_type,
                week=week,
            )

    if "contribution" in sections:
        generate_contribution_report(None, wb)

    if "fund" in sections:
        generate_organization_fund_report(year, month, report_type, wb, week=week)

    return wb
